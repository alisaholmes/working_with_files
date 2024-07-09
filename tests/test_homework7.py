import csv
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader

from script_os import ZIP_DIR


def test_pdf():
    with zipfile.ZipFile(ZIP_DIR) as zip_file:
        with zip_file.open('KPI.pdf') as pdf_file:
            reader = PdfReader(pdf_file)
            page = reader.pages[0]
            text = page.extract_text()
            assert 'Качественный показатель' in text


def test_xlsx():
    with zipfile.ZipFile(ZIP_DIR) as zip_f:
        with zip_f.open('characters.xlsx') as xlsx_file:
            workbook = load_workbook(xlsx_file)
            sheet = workbook.active
            value = sheet.cell(row=7, column=2).value
            name = 'Берест Владимир'
            assert name in value


def test_csv():
    with zipfile.ZipFile(ZIP_DIR) as zip_f:
        with zip_f.open('characters_story.csv') as csv_file:
            content = csv_file.read().decode('utf-8-sig')
            csvreader = list(csv.reader(content.splitlines()))
            second_row = csvreader[1]
            assert second_row[0] == 'Беликов Юрий;14.06.2001;глава 1'
